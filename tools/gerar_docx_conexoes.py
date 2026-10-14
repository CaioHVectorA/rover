import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def set_cell_padding(cell, top=120, bottom=120, left=160, right=160):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for margin_name, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{margin_name}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="D3D3D3", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def build_document():
    doc = docx.Document()

    # 1. Configuração de Margens (2.0 cm)
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Cores do tema
    PRIMARY = RGBColor(15, 76, 129)     # Azul Petróleo / Navy (#0F4C81)
    SECONDARY = RGBColor(41, 128, 185)  # Azul Médio (#2980B9)
    TEXT_DARK = RGBColor(44, 62, 80)    # Cinza Escuro (#2C3E50)
    HEX_HEADER = "0F4C81"
    HEX_ZEBRA = "F7F9FB"
    HEX_WHITE = "FFFFFF"

    # Estilo Normal
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Calibri'
    style_normal.font.size = Pt(10.5)
    style_normal.font.color.rgb = TEXT_DARK

    # Título Principal
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(4)
    run_title = p_title.add_run("ROVER MÓVEL ESP32-S3 & ARDUINO NANO")
    run_title.font.name = 'Arial'
    run_title.font.size = Pt(20)
    run_title.font.bold = True
    run_title.font.color.rgb = PRIMARY

    # Subtítulo
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(16)
    run_sub = p_sub.add_run("Projeto de Soldagem na Placa Perfurada: Componentes & Conexões")
    run_sub.font.name = 'Calibri'
    run_sub.font.size = Pt(13)
    run_sub.font.color.rgb = SECONDARY

    # Box de Resumo Executivo
    p_box = doc.add_paragraph()
    p_box.paragraph_format.space_after = Pt(14)
    r_box = p_box.add_run(
        "Este documento consolida de forma estrita e simplificada todos os componentes físicos e "
        "o mapeamento pino a pino para montagem da eletrônica interna em placa perfurada. "
        "A placa ESP32-S3-CAM fica alojada externamente na cabeça móvel Pan/Tilt, compartilhando "
        "apenas um chicote flexível de 3 fios (+5V, GND e TX) com a placa do chassi."
    )
    r_box.font.italic = True
    r_box.font.size = Pt(10)

    # -------------------------------------------------------------------------
    # SEÇÃO 1: LISTA DE COMPONENTES
    # -------------------------------------------------------------------------
    h1 = doc.add_heading(level=1)
    r_h1 = h1.add_run("1. Lista de Componentes Físicos")
    r_h1.font.name = 'Arial'
    r_h1.font.color.rgb = PRIMARY
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(6)

    # Tabela 1: Componentes Internos
    p_tab1 = doc.add_paragraph()
    p_tab1.paragraph_format.space_after = Pt(4)
    r_t1 = p_tab1.add_run("A. Módulos e Componentes na Placa Perfurada (Chassi)")
    r_t1.font.bold = True
    r_t1.font.color.rgb = PRIMARY

    comp_internos = [
        ("Arduino Nano V3 (ATmega328P)", "1 un", "Microcontrolador do chassi (controla TB6612, servos e lê bateria)"),
        ("Módulo Step-Down Mini 560 (5V / até 5A)", "1 un", "Regulador chaveado rebaixador de 7.4V para 5.0V estáveis"),
        ("Driver Duplo MOSFET TB6612FNG", "1 un", "Ponte H de tração para os 4 micromotores DC"),
        ("Transistor NPN BC337 (TO-92)", "1 un", "Inversor de sinal serial do ESP32 TX para Nano RX (D2)"),
        ("Resistor 1 kΩ (1/4W)", "1 un", "Resistor de base do transistor BC337"),
        ("Resistor 10 kΩ (1/4W)", "1 un", "Resistor de pull-up do coletor do BC337 ao +5V"),
        ("Barras de Pinos Fêmea (Soquetes 2.54 mm)", "2 pares", "Para encaixar o Nano e TB6612 sem soldar as placas diretamente"),
        ("Barras de Pinos Macho (Headers 2.54 mm)", "1 tira", "Para conexão dos servos (3 pinos) e chicote do ESP32 (3 pinos)"),
        ("Bornes de Parafuso KRE 2 vias (opcional)", "3 un", "Entrada da bateria (1x) e saídas dos motores esq/dir (2x)")
    ]

    table1 = doc.add_table(rows=len(comp_internos) + 1, cols=3)
    table1.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table1)

    headers1 = ["Componente", "Qtd", "Função no Circuito"]
    col_widths1 = [Inches(2.5), Inches(0.8), Inches(3.5)]

    # Cabeçalho Tabela 1
    hdr_cells1 = table1.rows[0].cells
    for i, title in enumerate(headers1):
        hdr_cells1[i].text = title
        set_cell_background(hdr_cells1[i], HEX_HEADER)
        set_cell_padding(hdr_cells1[i], top=100, bottom=100)
        p = hdr_cells1[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 1 else WD_ALIGN_PARAGRAPH.LEFT
        for run in p.runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9.5)

    # Linhas Tabela 1
    for r_idx, (c_name, c_qtd, c_fun) in enumerate(comp_internos):
        row_cells = table1.rows[r_idx + 1].cells
        row_cells[0].text = c_name
        row_cells[1].text = c_qtd
        row_cells[2].text = c_fun

        bg = HEX_ZEBRA if r_idx % 2 == 1 else HEX_WHITE
        for c_idx in range(3):
            set_cell_background(row_cells[c_idx], bg)
            set_cell_padding(row_cells[c_idx], top=80, bottom=80)
            p = row_cells[c_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx == 1 else WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.size = Pt(9)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Tabela 2: Componentes Externos
    p_tab2 = doc.add_paragraph()
    p_tab2.paragraph_format.space_after = Pt(4)
    r_t2 = p_tab2.add_run("B. Componentes Externos (Fora da Placa Perfurada)")
    r_t2.font.bold = True
    r_t2.font.color.rgb = PRIMARY

    comp_externos = [
        ("ESP32-S3-CAM N16R8 (Câmera 5MP)", "1 un", "Cabeça Pan/Tilt: Wi-Fi, streaming MJPEG e WebSocket (Conecta por 3 fios)"),
        ("Pack Bateria Li-Ion 2S (7.4V / 2500mAh)", "1 un", "Alimentação geral do sistema com placa BMS integrada"),
        ("Chave Gangorra KCD1-101 (2 Terminais)", "1 un", "Interruptor geral liga/desliga (corta o polo positivo da bateria)"),
        ("Motores DC JGA16-050 com Redução (16 mm)", "4 un", "Tração 4WD: 2 motores esquerdos e 2 direitos ligados em paralelo"),
        ("Micro Servos SG90 9g", "2 un", "Articulação Pan (Horizontal) e Tilt (Vertical) da câmera"),
        ("Rodas StickyMAX S20 22 mm (RoboCore)", "4 un", "Com cubo de fixação em alumínio CNC 6351-T6 incluso de fábrica")
    ]

    table2 = doc.add_table(rows=len(comp_externos) + 1, cols=3)
    table2.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table2)

    hdr_cells2 = table2.rows[0].cells
    for i, title in enumerate(headers1):
        hdr_cells2[i].text = title
        set_cell_background(hdr_cells2[i], HEX_HEADER)
        set_cell_padding(hdr_cells2[i], top=100, bottom=100)
        p = hdr_cells2[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 1 else WD_ALIGN_PARAGRAPH.LEFT
        for run in p.runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9.5)

    for r_idx, (c_name, c_qtd, c_fun) in enumerate(comp_externos):
        row_cells = table2.rows[r_idx + 1].cells
        row_cells[0].text = c_name
        row_cells[1].text = c_qtd
        row_cells[2].text = c_fun

        bg = HEX_ZEBRA if r_idx % 2 == 1 else HEX_WHITE
        for c_idx in range(3):
            set_cell_background(row_cells[c_idx], bg)
            set_cell_padding(row_cells[c_idx], top=80, bottom=80)
            p = row_cells[c_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx == 1 else WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.size = Pt(9)

    doc.add_page_break()

    # -------------------------------------------------------------------------
    # SEÇÃO 2: ESQUEMA VISUAL E ARQUITETURA
    # -------------------------------------------------------------------------
    h2 = doc.add_heading(level=1)
    r_h2 = h2.add_run("2. Esquema Elétrico e Diagramas Visuais")
    r_h2.font.name = 'Arial'
    r_h2.font.color.rgb = PRIMARY
    h2.paragraph_format.space_before = Pt(8)
    h2.paragraph_format.space_after = Pt(6)

    # Imagem 1: Esquema Elétrico KiCad
    p_img1_tit = doc.add_paragraph()
    r_i1 = p_img1_tit.add_run("Figura 1: Esquema Elétrico Completo de Interligação (KiCad)")
    r_i1.font.bold = True
    r_i1.font.size = Pt(10)

    path_kicad_sch = "docs/hardware/kicad/rover_perfurada_90x150/rover_conexoes.png"
    if os.path.exists(path_kicad_sch):
        p_img1 = doc.add_paragraph()
        p_img1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img1.paragraph_format.space_after = Pt(12)
        doc.add_picture(path_kicad_sch, width=Inches(6.4))
        p_cap1 = doc.add_paragraph()
        p_cap1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_c1 = p_cap1.add_run("Diagrama unifilar das conexões: Bateria, Mini 560, Nano, TB6612, Servos e Chicote ESP32.")
        r_c1.font.italic = True
        r_c1.font.size = Pt(8.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Imagem 2: Chicote ESP32 vs Chassi
    path_chicote = "docs/hardware/esquema_eletrico_chicote_esp32_nano.png"
    if os.path.exists(path_chicote):
        p_img2_tit = doc.add_paragraph()
        r_i2 = p_img2_tit.add_run("Figura 2: Arquitetura Distribuída (Cabeça Móvel vs Chassi Fixo)")
        r_i2.font.bold = True
        r_i2.font.size = Pt(10)

        p_img2 = doc.add_paragraph()
        p_img2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_picture(path_chicote, width=Inches(5.8))
        p_cap2 = doc.add_paragraph()
        p_cap2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_c2 = p_cap2.add_run("O chicote de 3 fios elimina a torção mecânica na movimentação Pan/Tilt.")
        r_c2.font.italic = True
        r_c2.font.size = Pt(8.5)

    doc.add_page_break()

    # -------------------------------------------------------------------------
    # SEÇÃO 3: MAPEAMENTO DE CONEXÕES PINO A PINO
    # -------------------------------------------------------------------------
    h3 = doc.add_heading(level=1)
    r_h3 = h3.add_run("3. Mapeamento Pino a Pino das Conexões")
    r_h3.font.name = 'Arial'
    r_h3.font.color.rgb = PRIMARY
    h3.paragraph_format.space_before = Pt(8)
    h3.paragraph_format.space_after = Pt(6)

    # 3.1 Alimentação
    p_sec1 = doc.add_paragraph()
    r_s1 = p_sec1.add_run("A. Distribuição de Alimentação e Terra")
    r_s1.font.bold = True
    r_s1.font.color.rgb = PRIMARY

    conexoes_alim = [
        ("Bateria Li-Ion (+)", "Entrada da Chave KCD1", "VBAT Bruto (7.4V nominal / 8.4V pico)"),
        ("Saída da Chave KCD1", "Step-Down Mini 560 (IN+) e TB6612 (VM)", "VBAT Chaveado para potência"),
        ("Bateria Li-Ion (-)", "Trilha Barramento GND", "Referência 0V Geral"),
        ("Mini 560 (IN- e OUT-)", "Trilha Barramento GND", "GND de entrada e saída do regulador"),
        ("Mini 560 (OUT+)", "Trilha Barramento +5V", "Fornece 5.0V estáveis (até 5A) para toda a lógica"),
        ("Trilha Barramento +5V", "Nano (5V), TB6612 (VCC), ESP32 (5V), Servos (5V), R2 (10k)", "Distribuição de 5V na placa"),
        ("Trilha Barramento GND", "Nano (GND), TB6612 (GND), ESP32 (GND), Servos (GND), BC337 (E)", "Distribuição de GND na placa")
    ]

    t_alim = doc.add_table(rows=len(conexoes_alim) + 1, cols=3)
    t_alim.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_alim)

    hdr_alim = t_alim.rows[0].cells
    for i, title in enumerate(["Ponto de Origem", "Ponto de Destino", "Descrição do Sinal"]):
        hdr_alim[i].text = title
        set_cell_background(hdr_alim[i], HEX_HEADER)
        set_cell_padding(hdr_alim[i], top=90, bottom=90)
        p = hdr_alim[i].paragraphs[0]
        for run in p.runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9)

    for r_idx, (o, d, s) in enumerate(conexoes_alim):
        row = t_alim.rows[r_idx + 1].cells
        row[0].text = o
        row[1].text = d
        row[2].text = s
        bg = HEX_ZEBRA if r_idx % 2 == 1 else HEX_WHITE
        for c_idx in range(3):
            set_cell_background(row[c_idx], bg)
            set_cell_padding(row[c_idx], top=70, bottom=70)
            p = row[c_idx].paragraphs[0]
            for run in p.runs:
                run.font.size = Pt(8.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 3.2 Comunicação Serial ESP32 ➔ Inversor BC337 ➔ Nano
    p_sec2 = doc.add_paragraph()
    r_s2 = p_sec2.add_run("B. Comunicação Serial (Chicote ESP32 de 3 Fios & Circuito Inversor)")
    r_s2.font.bold = True
    r_s2.font.color.rgb = PRIMARY

    conexoes_serial = [
        ("ESP32 Pino 5V (Externo)", "Trilha Barramento +5V", "Fio 1 do chicote: Alimentação da placa de vídeo"),
        ("ESP32 Pino GND (Externo)", "Trilha Barramento GND", "Fio 2 do chicote: Terra de sinal e alimentação"),
        ("ESP32 Pino IO43 TX (Externo)", "Resistor R1 (1 kΩ) - Terminal 1", "Fio 3 do chicote: Transmissão de comandos (3.3V)"),
        ("Resistor R1 (1 kΩ) - Terminal 2", "Transistor BC337 - Base (Pino 2)", "Entrada de corrente para chaveamento"),
        ("Transistor BC337 - Emissor (Pino 3)", "Trilha Barramento GND", "Referência de comutação"),
        ("Transistor BC337 - Coletor (Pino 1)", "Arduino Nano Pino D2 (RX)", "Sinal serial convertido e invertido para 5V"),
        ("Transistor BC337 - Coletor (Pino 1)", "Resistor R2 (10 kΩ) - Terminal 1", "Ponto de pull-up"),
        ("Resistor R2 (10 kΩ) - Terminal 2", "Trilha Barramento +5V", "Pull-up para garantir nível HIGH em 5.0V")
    ]

    t_ser = doc.add_table(rows=len(conexoes_serial) + 1, cols=3)
    t_ser.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_ser)

    hdr_ser = t_ser.rows[0].cells
    for i, title in enumerate(["Ponto de Origem", "Ponto de Destino", "Descrição do Sinal"]):
        hdr_ser[i].text = title
        set_cell_background(hdr_ser[i], HEX_HEADER)
        set_cell_padding(hdr_ser[i], top=90, bottom=90)
        p = hdr_ser[i].paragraphs[0]
        for run in p.runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9)

    for r_idx, (o, d, s) in enumerate(conexoes_serial):
        row = t_ser.rows[r_idx + 1].cells
        row[0].text = o
        row[1].text = d
        row[2].text = s
        bg = HEX_ZEBRA if r_idx % 2 == 1 else HEX_WHITE
        for c_idx in range(3):
            set_cell_background(row[c_idx], bg)
            set_cell_padding(row[c_idx], top=70, bottom=70)
            p = row[c_idx].paragraphs[0]
            for run in p.runs:
                run.font.size = Pt(8.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 3.3 Tração TB6612 e Motores
    p_sec3 = doc.add_paragraph()
    r_s3 = p_sec3.add_run("C. Tração 4WD: Arduino Nano ➔ TB6612FNG ➔ Motores JGA16-050")
    r_s3.font.bold = True
    r_s3.font.color.rgb = PRIMARY

    conexoes_motores = [
        ("Arduino Nano D12", "TB6612 Pino STBY", "Habilitação da Ponte H (HIGH = Motores ativos)"),
        ("Arduino Nano D3 (PWM)", "TB6612 Pino PWMA", "Velocidade analógica dos Motores Esquerdos"),
        ("Arduino Nano D4", "TB6612 Pino AIN1", "Sentido de rotação dos Motores Esquerdos"),
        ("Arduino Nano D5", "TB6612 Pino AIN2", "Sentido de rotação dos Motores Esquerdos"),
        ("Arduino Nano D6 (PWM)", "TB6612 Pino PWMB", "Velocidade analógica dos Motores Direitos"),
        ("Arduino Nano D7", "TB6612 Pino BIN1", "Sentido de rotação dos Motores Direitos"),
        ("Arduino Nano D8", "TB6612 Pino BIN2", "Sentido de rotação dos Motores Direitos"),
        ("TB6612 Saídas AO1 e AO2", "2x Motores Esquerdos (Paralelo)", "Potência Canal A (Frontal e Traseiro Esquerdos)"),
        ("TB6612 Saídas BO1 e BO2", "2x Motores Direitos (Paralelo)", "Potência Canal B (Frontal e Traseiro Direitos)")
    ]

    t_mot = doc.add_table(rows=len(conexoes_motores) + 1, cols=3)
    t_mot.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_mot)

    hdr_mot = t_mot.rows[0].cells
    for i, title in enumerate(["Ponto de Origem", "Ponto de Destino", "Descrição do Sinal"]):
        hdr_mot[i].text = title
        set_cell_background(hdr_mot[i], HEX_HEADER)
        set_cell_padding(hdr_mot[i], top=90, bottom=90)
        p = hdr_mot[i].paragraphs[0]
        for run in p.runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9)

    for r_idx, (o, d, s) in enumerate(conexoes_motores):
        row = t_mot.rows[r_idx + 1].cells
        row[0].text = o
        row[1].text = d
        row[2].text = s
        bg = HEX_ZEBRA if r_idx % 2 == 1 else HEX_WHITE
        for c_idx in range(3):
            set_cell_background(row[c_idx], bg)
            set_cell_padding(row[c_idx], top=70, bottom=70)
            p = row[c_idx].paragraphs[0]
            for run in p.runs:
                run.font.size = Pt(8.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 3.4 Servos Pan/Tilt
    p_sec4 = doc.add_paragraph()
    r_s4 = p_sec4.add_run("D. Servos da Câmera: Arduino Nano ➔ SG90")
    r_s4.font.bold = True
    r_s4.font.color.rgb = PRIMARY

    conexoes_servos = [
        ("Arduino Nano D9 (PWM)", "Header Servo Pan - Pino Sinal", "Controle angular Pan (Horizontal: 0° a 180°)"),
        ("Arduino Nano D10 (PWM)", "Header Servo Tilt - Pino Sinal", "Controle angular Tilt (Vertical: 0° a 180°)"),
        ("Trilha Barramento +5V", "Headers Servos - Pinos Centrais", "Alimentação dos servos (fio vermelho)"),
        ("Trilha Barramento GND", "Headers Servos - Pinos Terra", "Terra dos servos (fio marrom ou preto)")
    ]

    t_srv = doc.add_table(rows=len(conexoes_servos) + 1, cols=3)
    t_srv.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_srv)

    hdr_srv = t_srv.rows[0].cells
    for i, title in enumerate(["Ponto de Origem", "Ponto de Destino", "Descrição do Sinal"]):
        hdr_srv[i].text = title
        set_cell_background(hdr_srv[i], HEX_HEADER)
        set_cell_padding(hdr_srv[i], top=90, bottom=90)
        p = hdr_srv[i].paragraphs[0]
        for run in p.runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9)

    for r_idx, (o, d, s) in enumerate(conexoes_servos):
        row = t_srv.rows[r_idx + 1].cells
        row[0].text = o
        row[1].text = d
        row[2].text = s
        bg = HEX_ZEBRA if r_idx % 2 == 1 else HEX_WHITE
        for c_idx in range(3):
            set_cell_background(row[c_idx], bg)
            set_cell_padding(row[c_idx], top=70, bottom=70)
            p = row[c_idx].paragraphs[0]
            for run in p.runs:
                run.font.size = Pt(8.5)

    doc.add_page_break()

    # -------------------------------------------------------------------------
    # SEÇÃO 4: DICAS DE MONTAGEM E LAYOUT DA PLACA PERFURADA
    # -------------------------------------------------------------------------
    h4 = doc.add_heading(level=1)
    r_h4 = h4.add_run("4. Dicas Práticas de Soldagem e Layout")
    r_h4.font.name = 'Arial'
    r_h4.font.color.rgb = PRIMARY
    h4.paragraph_format.space_before = Pt(8)
    h4.paragraph_format.space_after = Pt(8)

    dicas = [
        ("Uso de Soquetes (Barras Fêmea):", "Nunca solde o Arduino Nano ou a placa TB6612FNG diretamente na placa perfurada. Solde barras de pinos fêmea no passo de 2.54 mm e encaixe os módulos como chips. Isso protege os componentes contra calor excessivo e permite substituição imediata ou gravação fora do robô se necessário."),
        ("Barramentos Mestres de Alimentação:", "Em placas de ilhas isoladas, crie duas trilhas contínuas de fio rígido estanhado (ou pontes de solda) atravessando a placa: uma para GND e outra para +5V. Isso permite puxar alimentação para qualquer módulo a partir da ilha vizinha, sem cruzar fios longos pelo meio da placa."),
        ("Conectores de Borda para Cabos Externos:", "Solde headers macho de 2.54 mm nas bordas da placa perfurada. Para os servos, monte blocos de 3 pinos [SINAL | 5V | GND] onde o plugue padrão dos servos encaixa diretamente sem cortes. Para o chicote do ESP32, monte um conector de 3 pinos [5V | GND | TX]."),
        ("Traçado dos Fios de Potência:", "Os fios que ligam o Step-Down Mini 560 e as saídas dos motores (TB6612 para os bornes) devem ser de bitola adequada (AWG 22 a 24) e o mais curtos possível para minimizar queda de tensão.")
    ]

    for d_tit, d_desc in dicas:
        p_d = doc.add_paragraph()
        p_d.paragraph_format.space_after = Pt(6)
        r_dt = p_d.add_run(f"• {d_tit} ")
        r_dt.font.bold = True
        r_dt.font.color.rgb = PRIMARY
        r_dd = p_d.add_run(d_desc)
        r_dd.font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Imagem 3: Layout sugerido na perfurada
    path_layout = "docs/hardware/kicad/rover_perfurada_90x150/montagem_superior.png"
    if os.path.exists(path_layout):
        p_img3_tit = doc.add_paragraph()
        r_i3 = p_img3_tit.add_run("Figura 3: Sugestão de Layout Físico na Placa Perfurada (90 x 150 mm)")
        r_i3.font.bold = True
        r_i3.font.size = Pt(10)

        p_img3 = doc.add_paragraph()
        p_img3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_picture(path_layout, width=Inches(4.2))
        p_cap3 = doc.add_paragraph()
        p_cap3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_c3 = p_cap3.add_run("Disposição das ilhas e posicionamento relativo dos módulos e componentes.")
        r_c3.font.italic = True
        r_c3.font.size = Pt(8.5)

    # Salva o arquivo .docx
    output_docx = "docs/hardware/PROJETO_CONEXOES_PLACA_PERFURADA.docx"
    doc.save(output_docx)
    print(f"Documento DOCX salvo com sucesso em: {output_docx}")

if __name__ == "__main__":
    build_document()
