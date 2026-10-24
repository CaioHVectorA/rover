import openpyxl

wb = openpyxl.load_workbook(r'docs\hardware\compras_rover_reorganizado.xlsx')
ws = wb['UsinaInfo']
for r in [8, 14, 15]:
    cell_h = ws.cell(r, 8)
    print(f"Row {r} col 8 value: {cell_h.value}, hyperlink: {cell_h.hyperlink.target if cell_h.hyperlink else None}")
    for col in range(2, 12):
        c = ws.cell(r, col)
        print(f"  Col {col} ({c.coordinate}): val={c.value}, num_format={c.number_format}, font={c.font.name if c.font else None}")
