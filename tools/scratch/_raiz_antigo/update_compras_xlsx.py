import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.hyperlink import Hyperlink
import copy
import shutil

src_wb_path = r'docs\hardware\compras_rover_reorganizado.xlsx'
tgt_wb_path = r'docs\hardware\compras_rover_reorganizado_atualizado.xlsx'
wb = openpyxl.load_workbook(src_wb_path)

# ==========================================
# 1. Update UsinaInfo Sheet
# ==========================================
ws_u = wb['UsinaInfo']

# Row 16 will be the new item. Let's insert row at 16
ws_u.insert_rows(16)

# Copy styles from row 15
for col in range(1, ws_u.max_column + 1):
    src_cell = ws_u.cell(15, col)
    tgt_cell = ws_u.cell(16, col)
    if src_cell.has_style:
        tgt_cell.font = copy.copy(src_cell.font)
        tgt_cell.border = copy.copy(src_cell.border)
        tgt_cell.fill = copy.copy(src_cell.fill)
        tgt_cell.number_format = copy.copy(src_cell.number_format)
        tgt_cell.protection = copy.copy(src_cell.protection)
        tgt_cell.alignment = copy.copy(src_cell.alignment)

# Set values for row 16: Conversor de nível lógico
ws_u.cell(16, 2, '03741')
ws_u.cell(16, 3, 'Conversor de nível lógico 3,3V / 5V bidirecional — CNL35')
ws_u.cell(16, 4, 1)
ws_u.cell(16, 5, 7.50)
ws_u.cell(16, 6, '=ROUND(D16*E16,2)')
ws_u.cell(16, 7, '=ROUND(F16*(1-$D$5),2)')

link_u = 'https://www.usinainfo.com.br/conversores-de-sinal/conversor-de-nivel-logico-33-5v-bidirecional-cnl35-3741.html'
cell_link_u = ws_u.cell(16, 8, 'Abrir produto')
cell_link_u.hyperlink = Hyperlink(ref=cell_link_u.coordinate, target=link_u)
cell_link_u.font = Font(name='Segoe UI', size=9, color='0563C1', underline='single')

ws_u.cell(16, 9, 'Alta')
ws_u.cell(16, 10, 'Em estoque')
ws_u.cell(16, 11, 'Adaptação 3,3 V (ESP32) ⇄ 5 V (Nano) para UART nativa a 115200 baud, eliminando o transistor BC337.')

# Format money columns
ws_u.cell(16, 5).number_format = '"R$ "#,##0.00'
ws_u.cell(16, 6).number_format = '"R$ "#,##0.00'
ws_u.cell(16, 7).number_format = '"R$ "#,##0.00'

# Update Totals on UsinaInfo (now rows 19 to 22)
ws_u.cell(19, 6, '=SUM(F8:F16)')
ws_u.cell(19, 7, '=SUM(G8:G16)')
ws_u.cell(20, 6, '=SUMIF(I8:I16,"Alta",F8:F16)')
ws_u.cell(20, 7, '=SUMIF(I8:I16,"Alta",G8:G16)')
ws_u.cell(21, 6, '=SUMIF(I8:I16,"Média",F8:F16)')
ws_u.cell(21, 7, '=SUMIF(I8:I16,"Média",G8:G16)')
ws_u.cell(22, 6, '=SUMIF(I8:I16,"Baixa",F8:F16)')
ws_u.cell(22, 7, '=SUMIF(I8:I16,"Baixa",G8:G16)')

# Update Data Validation for UsinaInfo
for dv in ws_u.data_validations.dataValidation:
    dv.sqref = 'I8:I16'


# ==========================================
# 2. Update Mercado Livre Sheet
# ==========================================
ws_ml = wb['Mercado Livre']

# Row 11 will be the new item. Let's insert row at 11
ws_ml.insert_rows(11)

# Copy styles from row 10
for col in range(1, ws_ml.max_column + 1):
    src_cell = ws_ml.cell(10, col)
    tgt_cell = ws_ml.cell(11, col)
    if src_cell.has_style:
        tgt_cell.font = copy.copy(src_cell.font)
        tgt_cell.border = copy.copy(src_cell.border)
        tgt_cell.fill = copy.copy(src_cell.fill)
        tgt_cell.number_format = copy.copy(src_cell.number_format)
        tgt_cell.protection = copy.copy(src_cell.protection)
        tgt_cell.alignment = copy.copy(src_cell.alignment)

# Set values for row 11: Áudio I2S (MAX98357A + Alto-falante 3W)
ws_ml.cell(11, 2, 'ML-MAX98357A')
ws_ml.cell(11, 3, 'Módulo amplificador I2S MAX98357A + Mini alto-falante 3W')
ws_ml.cell(11, 4, 1)
ws_ml.cell(11, 5, 25.00)
ws_ml.cell(11, 6, '=ROUND(D11*E11,2)')
ws_ml.cell(11, 7, '=ROUND(F11*(1-$D$5),2)')

link_ml = 'https://lista.mercadolivre.com.br/modulo-max98357a'
cell_link_ml = ws_ml.cell(11, 8, 'Revisar anúncio')
cell_link_ml.hyperlink = Hyperlink(ref=cell_link_ml.coordinate, target=link_ml)
cell_link_ml.font = Font(name='Segoe UI', size=9, color='0563C1', underline='single')

ws_ml.cell(11, 9, 'Média')
ws_ml.cell(11, 10, 'A recotar Full')
ws_ml.cell(11, 11, 'DAC + amplificador Classe D 3,2 W via I2S nativo do ESP32 (3 pinos) para buzina real, efeitos sonoros e alertas.')

# Format money columns
ws_ml.cell(11, 5).number_format = '"R$ "#,##0.00'
ws_ml.cell(11, 6).number_format = '"R$ "#,##0.00'
ws_ml.cell(11, 7).number_format = '"R$ "#,##0.00'

# Update Totals on Mercado Livre
# Row 14 is Total
ws_ml.cell(14, 6, '=SUM(F8:F11)')
ws_ml.cell(14, 7, '=SUM(G8:G11)')

# Insert a row for Subtotal Média
ws_ml.insert_rows(15)
for col in range(1, ws_ml.max_column + 1):
    src = ws_ml.cell(16, col)
    tgt = ws_ml.cell(15, col)
    if src.has_style:
        tgt.font = copy.copy(src.font)
        tgt.border = copy.copy(src.border)
        tgt.fill = copy.copy(src.fill)
        tgt.number_format = copy.copy(src.number_format)
        tgt.alignment = copy.copy(src.alignment)

ws_ml.cell(15, 3, 'Subtotal prioridade média (estimado)')
ws_ml.cell(15, 6, '=SUMIF(I8:I11,"Média",F8:F11)')
ws_ml.cell(15, 7, '=SUMIF(I8:I11,"Média",G8:G11)')
ws_ml.cell(15, 6).number_format = '"R$ "#,##0.00'
ws_ml.cell(15, 7).number_format = '"R$ "#,##0.00'

ws_ml.cell(16, 3, 'Subtotal prioridade baixa (estimado)')
ws_ml.cell(16, 6, '=SUMIF(I8:I11,"Baixa",F8:F11)')
ws_ml.cell(16, 7, '=SUMIF(I8:I11,"Baixa",G8:G11)')
ws_ml.cell(16, 6).number_format = '"R$ "#,##0.00'
ws_ml.cell(16, 7).number_format = '"R$ "#,##0.00'

# Update Data Validation for Mercado Livre
for dv in ws_ml.data_validations.dataValidation:
    dv.sqref = 'I8:I11'


# ==========================================
# 3. Update Resumo Geral Sheet
# ==========================================
ws_rg = wb['Resumo Geral']

# Row 6: Cards
ws_rg.cell(6, 2, "='UsinaInfo'!F19+'Mercado Livre'!F14+'Curto Circuito'!F10")
ws_rg.cell(6, 4, "='UsinaInfo'!G19+'Mercado Livre'!G14+'Curto Circuito'!G10")

# Row 11: UsinaInfo
ws_rg.cell(11, 5, "='UsinaInfo'!F19")
ws_rg.cell(11, 6, "='UsinaInfo'!G19")

# Row 12: Mercado Livre
ws_rg.cell(12, 5, "='Mercado Livre'!F14")
ws_rg.cell(12, 6, "='Mercado Livre'!G14")

# Row 19: Alta
ws_rg.cell(19, 4, "=SUMIF(UsinaInfo!I8:I16,\"Alta\",UsinaInfo!F8:F16)+SUMIF('Mercado Livre'!I8:I11,\"Alta\",'Mercado Livre'!F8:F11)+SUMIF('Curto Circuito'!I8:I8,\"Alta\",'Curto Circuito'!F8:F8)")
ws_rg.cell(19, 5, "=SUMIF(UsinaInfo!I8:I16,\"Alta\",UsinaInfo!G8:G16)+SUMIF('Mercado Livre'!I8:I11,\"Alta\",'Mercado Livre'!G8:G11)+SUMIF('Curto Circuito'!I8:I8,\"Alta\",'Curto Circuito'!G8:G8)")

# Row 20: Média
ws_rg.cell(20, 4, "=SUMIF(UsinaInfo!I8:I16,\"Média\",UsinaInfo!F8:F16)+SUMIF('Mercado Livre'!I8:I11,\"Média\",'Mercado Livre'!F8:F11)+SUMIF('Curto Circuito'!I8:I8,\"Média\",'Curto Circuito'!F8:F8)")
ws_rg.cell(20, 5, "=SUMIF(UsinaInfo!I8:I16,\"Média\",UsinaInfo!G8:G16)+SUMIF('Mercado Livre'!I8:I11,\"Média\",'Mercado Livre'!G8:G11)+SUMIF('Curto Circuito'!I8:I8,\"Média\",'Curto Circuito'!G8:G8)")

# Row 21: Baixa
ws_rg.cell(21, 4, "=SUMIF(UsinaInfo!I8:I16,\"Baixa\",UsinaInfo!F8:F16)+SUMIF('Mercado Livre'!I8:I11,\"Baixa\",'Mercado Livre'!F8:F11)+SUMIF('Curto Circuito'!I8:I8,\"Baixa\",'Curto Circuito'!F8:F8)")
ws_rg.cell(21, 5, "=SUMIF(UsinaInfo!I8:I16,\"Baixa\",UsinaInfo!G8:G16)+SUMIF('Mercado Livre'!I8:I11,\"Baixa\",'Mercado Livre'!G8:G11)+SUMIF('Curto Circuito'!I8:I8,\"Baixa\",'Curto Circuito'!G8:G8)")

# Save to updated workbook
wb.save(tgt_wb_path)
print("Successfully generated:", tgt_wb_path)
