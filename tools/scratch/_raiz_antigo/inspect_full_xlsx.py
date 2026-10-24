import openpyxl

wb = openpyxl.load_workbook(r'docs\hardware\compras_rover_reorganizado.xlsx')
for sheet in wb.sheetnames:
    ws = wb[sheet]
    print(f"=== Sheet: {sheet} (max_row={ws.max_row}, max_column={ws.max_column}) ===")
    for r in range(1, ws.max_row + 1):
        row_vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        if any(v is not None for v in row_vals):
            print(f"Row {r:2d}: {row_vals}")
