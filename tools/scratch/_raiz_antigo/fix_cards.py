import openpyxl

wb = openpyxl.load_workbook(r'docs\hardware\compras_rover_reorganizado_atualizado.xlsx')
ws = wb['Lista de Compras']

# Row 5 cards point to row 36
ws.cell(5, 1, '=E36')
ws.cell(5, 3, '=D36')
ws.cell(5, 5, '=C5-A5')

wb.save(r'docs\hardware\compras_rover_reorganizado_atualizado.xlsx')
print('Cards updated to point to E36 and D36!')
