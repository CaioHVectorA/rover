import openpyxl, sys

sys.stdout.reconfigure(encoding='utf-8')
wb = openpyxl.load_workbook(r'docs\hardware\compras_rover_reorganizado_atualizado.xlsx')
for sheetname in ['Resumo Geral', 'UsinaInfo', 'Mercado Livre']:
    ws = wb[sheetname]
    print(f"\n=================== {sheetname} ===================")
    for r in range(1, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        if any(v is not None for v in vals):
            print(f"R{r:02d}: {vals}")
