import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.hyperlink import Hyperlink
import copy

wb_path = r'docs\hardware\compras_rover_reorganizado_atualizado.xlsx'
wb = openpyxl.load_workbook(wb_path)

ws_u = wb['UsinaInfo']

# Current items in UsinaInfo: rows 8 to 16 (9 items).
# Row 12 is the Antenna. Let's make its product name super clear that it includes the pigtail cable:
ws_u.cell(12, 3, 'Kit Antena Wi-Fi 2,4GHz 3dBi + Cabo Pigtail U.FL (IPEX1) para SMA Fêmea')
ws_u.cell(12, 11, 'Kit completo: inclui a antena de 3dBi e o cabo adaptador Pigtail U.FL para rosca SMA fêmea.')

# Now let's insert 2 new rows at row 17 and row 18 for:
# 1. Display TFT 3,5" SPI Touch (09989)
# 2. Push Button Chave Táctil 6x6x5mm (Kit com 5 un - Qtd: 2 = 10 botões) (02989)
ws_u.insert_rows(17, amount=2)

for new_row in [17, 18]:
    for col in range(1, ws_u.max_column + 1):
        src_cell = ws_u.cell(16, col)
        tgt_cell = ws_u.cell(new_row, col)
        if src_cell.has_style:
            tgt_cell.font = copy.copy(src_cell.font)
            tgt_cell.border = copy.copy(src_cell.border)
            tgt_cell.fill = copy.copy(src_cell.fill)
            tgt_cell.number_format = copy.copy(src_cell.number_format)
            tgt_cell.protection = copy.copy(src_cell.protection)
            tgt_cell.alignment = copy.copy(src_cell.alignment)

# Row 17: Display TFT 3.5"
ws_u.cell(17, 2, '09989')
ws_u.cell(17, 3, 'Display TFT LCD 3,5" SPI Touch Screen com Slot SD (ILI9488)')
ws_u.cell(17, 4, 1)
ws_u.cell(17, 5, 119.60)
ws_u.cell(17, 6, '=ROUND(D17*E17,2)')
ws_u.cell(17, 7, '=ROUND(F17*(1-$D$5),2)')

link_tft35 = 'https://www.usinainfo.com.br/display-arduino/tft-lcd-35-spi-touch-screen-com-slot-sd-8570.html'
cell_tft35 = ws_u.cell(17, 8, 'Abrir produto')
cell_tft35.hyperlink = Hyperlink(ref=cell_tft35.coordinate, target=link_tft35)
cell_tft35.font = Font(name='Segoe UI', size=9, color='0563C1', underline='single')

ws_u.cell(17, 9, 'Alta')
ws_u.cell(17, 10, 'Em estoque')
ws_u.cell(17, 11, 'Display grande 480x320 com touch screen (XPT2046) e slot MicroSD, lógica nativa 3,3 V via SPI.')

ws_u.cell(17, 5).number_format = '"R$ "#,##0.00'
ws_u.cell(17, 6).number_format = '"R$ "#,##0.00'
ws_u.cell(17, 7).number_format = '"R$ "#,##0.00'

# Row 18: Botões Push Button (10 unidades = 2 kits de 5)
ws_u.cell(18, 2, '02989')
ws_u.cell(18, 3, 'Push Button Chave Táctil 6x6x5mm (2x Kits com 5 un = 10 botões)')
ws_u.cell(18, 4, 2)
ws_u.cell(18, 5, 2.46)
ws_u.cell(18, 6, '=ROUND(D18*E18,2)')
ws_u.cell(18, 7, '=ROUND(F18*(1-$D$5),2)')

link_btn = 'https://www.usinainfo.com.br/push-buttons/push-button-chave-tactil-6x6x5-kit-com-5-unidades-2989.html'
cell_btn = ws_u.cell(18, 8, 'Abrir produto')
cell_btn.hyperlink = Hyperlink(ref=cell_btn.coordinate, target=link_btn)
cell_btn.font = Font(name='Segoe UI', size=9, color='0563C1', underline='single')

ws_u.cell(18, 9, 'Alta')
ws_u.cell(18, 10, 'Em estoque')
ws_u.cell(18, 11, '10 chaves tácteis momentâneas para comandos físicos, reset, calibração e botões de interface.')

ws_u.cell(18, 5).number_format = '"R$ "#,##0.00'
ws_u.cell(18, 6).number_format = '"R$ "#,##0.00'
ws_u.cell(18, 7).number_format = '"R$ "#,##0.00'

# Update Totals on UsinaInfo (now rows 21 to 24)
ws_u.cell(21, 6, '=SUM(F8:F18)')
ws_u.cell(21, 7, '=SUM(G8:G18)')
ws_u.cell(22, 6, '=SUMIF(I8:I18,"Alta",F8:F18)')
ws_u.cell(22, 7, '=SUMIF(I8:I18,"Alta",G8:G18)')
ws_u.cell(23, 6, '=SUMIF(I8:I18,"Média",F8:F18)')
ws_u.cell(23, 7, '=SUMIF(I8:I18,"Média",G8:G18)')
ws_u.cell(24, 6, '=SUMIF(I8:I18,"Baixa",F8:F18)')
ws_u.cell(24, 7, '=SUMIF(I8:I18,"Baixa",G8:G18)')

for dv in ws_u.data_validations.dataValidation:
    dv.sqref = 'I8:I18'

# Update Resumo Geral
ws_rg = wb['Resumo Geral']

# Row 6: Cards
ws_rg.cell(6, 2, "='UsinaInfo'!F21+'Mercado Livre'!F14+'Curto Circuito'!F10")
ws_rg.cell(6, 4, "='UsinaInfo'!G21+'Mercado Livre'!G14+'Curto Circuito'!G10")

# Row 11: UsinaInfo
ws_rg.cell(11, 5, "='UsinaInfo'!F21")
ws_rg.cell(11, 6, "='UsinaInfo'!G21")

# Row 19: Alta
ws_rg.cell(19, 4, "=SUMIF(UsinaInfo!I8:I18,\"Alta\",UsinaInfo!F8:F18)+SUMIF('Mercado Livre'!I8:I11,\"Alta\",'Mercado Livre'!F8:F11)+SUMIF('Curto Circuito'!I8:I8,\"Alta\",'Curto Circuito'!F8:F8)")
ws_rg.cell(19, 5, "=SUMIF(UsinaInfo!I8:I18,\"Alta\",UsinaInfo!G8:G18)+SUMIF('Mercado Livre'!I8:I11,\"Alta\",'Mercado Livre'!G8:G11)+SUMIF('Curto Circuito'!I8:I8,\"Alta\",'Curto Circuito'!G8:G8)")

# Row 20: Média
ws_rg.cell(20, 4, "=SUMIF(UsinaInfo!I8:I18,\"Média\",UsinaInfo!F8:F18)+SUMIF('Mercado Livre'!I8:I11,\"Média\",'Mercado Livre'!F8:F11)+SUMIF('Curto Circuito'!I8:I8,\"Média\",'Curto Circuito'!F8:F8)")
ws_rg.cell(20, 5, "=SUMIF(UsinaInfo!I8:I18,\"Média\",UsinaInfo!G8:G18)+SUMIF('Mercado Livre'!I8:I11,\"Média\",'Mercado Livre'!G8:G11)+SUMIF('Curto Circuito'!I8:I8,\"Média\",'Curto Circuito'!G8:G8)")

# Row 21: Baixa
ws_rg.cell(21, 4, "=SUMIF(UsinaInfo!I8:I18,\"Baixa\",UsinaInfo!F8:F18)+SUMIF('Mercado Livre'!I8:I11,\"Baixa\",'Mercado Livre'!F8:F11)+SUMIF('Curto Circuito'!I8:I8,\"Baixa\",'Curto Circuito'!F8:F8)")
ws_rg.cell(21, 5, "=SUMIF(UsinaInfo!I8:I18,\"Baixa\",UsinaInfo!G8:G18)+SUMIF('Mercado Livre'!I8:I11,\"Baixa\",'Mercado Livre'!G8:G11)+SUMIF('Curto Circuito'!I8:I8,\"Baixa\",'Curto Circuito'!G8:G8)")

wb.save(wb_path)
print("Updated successfully:", wb_path)
